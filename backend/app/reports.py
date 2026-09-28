import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def _pdf_escape(value):
    return str(value).replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def executive_pdf(metrics, demands):
    lines = ["PAC DIGITAL MPRJ - RELATORIO EXECUTIVO", f"Exercicio 2026 | Emitido em {date.today().strftime('%d/%m/%Y')}", "", f"Demandas: {metrics['demands']}", f"Planejado: R$ {metrics['planned_value']:,.2f}", f"Executado: R$ {metrics['executed_value']:,.2f}", f"Taxa de execucao: {metrics['execution_rate']:.2f}%", f"Risco alto: {metrics['risk_demands']} | Risco medio: {metrics['medium_risk_demands']}", f"Alteracoes: {metrics['altered_demands']} | Extraordinarias: {metrics['extraordinary_inclusions']}", "", "CARTEIRA"]
    for item in demands:
        amount = float(item.adjusted_value or item.revised_value or item.original_value)
        lines.append(f"{item.code} | {item.title[:58]} | R$ {amount:,.2f} | {item.execution_status}")
    lines += ["", "Documento demonstrativo. Carteira ficticia; inteligencia PNCP baseada em dados publicos."]
    commands = ["BT", "/F1 15 Tf", "50 800 Td"]
    for index, line in enumerate(lines):
        if index == 1: commands += ["/F1 9 Tf"]
        commands += [f"({_pdf_escape(line)}) Tj", "0 -20 Td"]
    commands.append("ET"); stream = "\n".join(commands).encode("cp1252", errors="replace")
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>", b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>", b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    result = bytearray(b"%PDF-1.4\n"); offsets = [0]
    for number, obj in enumerate(objects, 1): offsets.append(len(result)); result += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(result); result += f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode()
    for offset in offsets[1:]: result += f"{offset:010d} 00000 n \n".encode()
    result += f"trailer << /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode(); return bytes(result)


def executive_xlsx(metrics, demands):
    book = Workbook(); summary = book.active; summary.title = "Resumo executivo"; wine = "7D2425"; gold = "C6A34F"
    summary.append(["PAC Digital MPRJ", "Relatório executivo demonstrativo"]); summary.append(["Indicador", "Valor"])
    for label, key in [("Demandas", "demands"), ("Planejado", "planned_value"), ("Executado", "executed_value"), ("Execução (%)", "execution_rate"), ("Risco alto", "risk_demands"), ("Alterações", "altered_demands")]: summary.append([label, metrics[key]])
    for cell in summary[1]: cell.fill = PatternFill("solid", fgColor=wine); cell.font = Font(color="FFFFFF", bold=True)
    for cell in summary[2]: cell.fill = PatternFill("solid", fgColor=gold); cell.font = Font(bold=True)
    sheet = book.create_sheet("Carteira"); sheet.append(["Código", "Objeto", "Unidade", "Categoria", "Planejado", "Executado", "Situação"])
    for cell in sheet[1]: cell.fill = PatternFill("solid", fgColor=wine); cell.font = Font(color="FFFFFF", bold=True)
    for item in demands: sheet.append([item.code, item.title, item.unit, item.category, float(item.adjusted_value or item.revised_value or item.original_value), float(item.executed_value), item.execution_status])
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions; stream = io.BytesIO(); book.save(stream); return stream.getvalue()
